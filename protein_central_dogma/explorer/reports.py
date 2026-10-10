from __future__ import annotations

import io
import json
import zipfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def safe_cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        if value.startswith(("=", "+", "-", "@")):
            value = "'" + value
        if len(value) > 30000:
            value = value[:29900] + " [TRUNCATED IN EXCEL; full value is in JSON/FASTA]"
    return value


def add_sheet(wb, name, rows, columns=None):
    ws = wb.create_sheet(name)
    columns = columns or list(dict.fromkeys(k for row in rows for k in row))
    if not columns:
        columns = ["Status"]
        rows = [{"Status": "No available records"}]
    ws.append(columns)
    for row in rows:
        ws.append([safe_cell(row.get(k)) for k in columns])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="123D50")
        c.font = Font(color="FFFFFF", bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        text = " ".join(str(c.value or "") for c in row).lower()
        colour = "E1F4E9" if "exact full-length reference" in text else "FFF3DA" if "unavailable" in text or "not verified" in text else "F3F7FA"
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=colour)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(55, max(16, len(str(col[0].value or "")) + 2))
    return ws


def workbook(queries, jobs, details):
    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Queries", [vars(q) for q in queries])
    candidate_rows = [{k: v for k, v in c.items() if k != "hsps"} for job in jobs for c in job.get("candidates", [])]
    add_sheet(wb, "Candidates", candidate_rows)
    summary, genes, genomes, retrievals, loci, transcripts, structures, translations, sources, issues = [], [], [], [], [], [], [], [], [], []
    for key, d in details.items():
        summary.append({"selection": key, "protein_accession": d["protein_accession"], "organism": d["organism"],
                        "taxid": d.get("taxid"), "gene_ids": d.get("gene_ids", []), "match": d["match"],
                        "gene_names": d.get("gene_names", []),
                        "reference_protein_length": len(d["protein_sequence"]), "retrieved_at": d["retrieved_at"],
                        "linked_nucleotide_ids": d.get("linked_nucleotide_ids", [])})
        for gene in d["genes"]:
            genes.append({"selection": key, **{k: v for k, v in gene.items() if k != "loci"}})
            loci.extend({"selection": key, **l} for l in gene["loci"])
        for genome in d.get("genomic", []):
            genomes.append({"selection": key, **{k: v for k, v in genome.items() if k not in ("sequence", "features")}, "sequence_length": len(genome["sequence"])})
            structures.extend({"selection": key, **f} for f in genome["features"])
        retrievals.extend({"selection": key, **r} for r in d.get("record_retrievals", []))
        for t in d["transcripts"]:
            transcripts.append({"selection": key, "accession": t["accession"], "molecule_type": t["molecule_type"],
                                "is_transcript": t["is_transcript"], "transcript_length": len(t["transcript_sequence"]),
                                "cds_length": len(t["translation"]["cds"]), "cds_location": t["cds_location"], "source": t["source"]})
            structures.extend({"selection": key, **f, "coordinate_context": "transcript or nucleotide record"} for f in t["features"] + t["utrs"])
            translations.append({"selection": key, "accession": t["accession"], **{k: v for k, v in t["translation"].items() if k not in ("cds", "computed_protein", "annotated_protein")}})
        sources.extend({"selection": key, **s, "report_generated": d["retrieved_at"]} for s in d["sources"])
        issues.extend({"selection": key, "warning": w} for w in d["warnings"])
    add_sheet(wb, "Protein_Details", summary)
    add_sheet(wb, "Protein_Features", [{"selection": key, **f} for key, d in details.items() for f in d.get("protein_features", [])])
    add_sheet(wb, "Genes", genes)
    add_sheet(wb, "Gene_Loci", loci)
    add_sheet(wb, "Genomic_Records", genomes)
    add_sheet(wb, "Transcripts_CDS", transcripts)
    add_sheet(wb, "Gene_Structure", structures)
    add_sheet(wb, "Translation_Checks", translations)
    add_sheet(wb, "Sources", sources)
    add_sheet(wb, "Retrieval_Dates", retrievals)
    add_sheet(wb, "Warnings", issues)
    add_sheet(wb, "Methods", [
        {"item": "Searches", "details": [{k: v for k, v in j.items() if k not in ("raw_blast", "candidates", "query")} for j in jobs]},
        {"item": "Identity", "details": "Identity is for the best scoring HSP, not a global alignment. Coverage is the union of all HSP coordinates."},
        {"item": "Exact match", "details": "Requires direct comparison of the full retrieved reference sequence; alignment scores alone are insufficient."},
        {"item": "Coordinates", "details": "All exported numeric coordinates are 1-based inclusive. GenBank location strings may use Biopython's 0-based representation."},
        {"item": "Availability", "details": "Capped database search; every returned candidate is preserved, but not every possible organism/gene is guaranteed."},
        {"item": "Biology", "details": "Reference links are not proof of sample origin. Protein cannot uniquely determine DNA. RNA-seq needs experimental data."},
        {"item": "Retrieval dates", "details": "Report retrieval time is recorded; cached record responses may be up to seven days old."},
    ])
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def fasta_text(entries):
    lines = []
    for identifier, seq in entries:
        if not seq:
            continue
        identifier = " ".join(str(identifier).splitlines())
        lines.append(">" + identifier)
        lines.extend(seq[i:i + 70] for i in range(0, len(seq), 70))
    return "\n".join(lines) + ("\n" if lines else "")


def fasta_files(details):
    proteins, query, transcripts, cds, computed, genomic = [], [], [], [], [], []
    for key, d in details.items():
        label = f"{key} reference={d['protein_accession']}"
        proteins.append((label, d["protein_sequence"]))
        query.append((f"{key} input_query", d["query_sequence"]))
        for t in d["transcripts"]:
            tlabel = f"{key} {t['accession']} protein={d['protein_accession']}"
            if t["is_transcript"] and t["transcript_sequence"]:
                transcripts.append((tlabel + " transcript_cDNA_alphabet", t["transcript_sequence"]))
            cds.append((tlabel + " CDS", t["translation"]["cds"]))
            computed.append((tlabel + " computed_translation", t["translation"]["computed_protein"]))
        for g in d.get("genomic", []):
            genomic.append((f"{key} {g['accession']}:{g['start']}-{g['end']} strand={g['strand']} orientation={g['orientation']}", g["sequence"]))
    return {name: fasta_text(entries) for name, entries in [
        ("query_proteins.fasta", query), ("reference_proteins.fasta", proteins),
        ("transcripts.fasta", transcripts), ("cds.fasta", cds),
        ("computed_proteins.fasta", computed), ("genomic_dna.fasta", genomic)]}


def structure_figure(detail):
    genome = detail.get("genomic", [])
    rows = []
    for g in genome:
        features = g["features"]
        parents = list(dict.fromkeys(f.get("parent", "gene") for f in features if f["feature"] != "gene"))
        for parent in parents:
            label = f"{g['accession']} / {parent}\norigin {g['start']:,}; strand {g['strand']}"
            rows.append((label, [f for f in features if f.get("parent", "gene") == parent], "genomic", g["start"]))
    if not rows:
        for t in detail["transcripts"]:
            rows.append((t["accession"], t["features"] + t["utrs"], "transcript" if t["is_transcript"] else "nucleotide record", 1))
    fig, ax = plt.subplots(figsize=(10, max(3, 0.8 * len(rows) + 1.7)))
    colours = {"CDS": "#168D89", "exon": "#8CCBB1", "mRNA": "#B9CCD7", "intron": "#657786", "5prime UTR": "#E6AB02", "3prime UTR": "#E6AB02", "5'UTR": "#E6AB02", "3'UTR": "#E6AB02"}
    if not rows:
        ax.text(0.5, 0.5, "Gene structure unavailable for this record", ha="center", transform=ax.transAxes)
        ax.set_axis_off()
    else:
        for index, (label, features, context, origin) in enumerate(rows):
            if not features:
                continue
            left, right = min(f["start"] for f in features) - origin + 1, max(f["end"] for f in features) - origin + 1
            ax.plot([left, right], [index, index], color="#A9BAC4", lw=1)
            for feature in sorted(features, key=lambda f: f["feature"] == "CDS"):
                width = feature["end"] - feature["start"] + 1
                height = 0.52 if feature["feature"] == "CDS" else 0.26
                ax.add_patch(Rectangle((feature["start"] - origin + 1, index - height / 2), width, height, color=colours.get(feature["feature"], "#B9CCD7")))
        ax.set_yticks(range(len(rows)), [row[0] for row in rows], fontsize=8)
        ax.set_ylim(-0.7, len(rows) - 0.3)
        ax.autoscale_view(scalex=True, scaley=False)
        ax.set_xlabel("Position within each reference region (bp); absolute origin shown in each row label")
        ax.ticklabel_format(axis="x", style="plain", useOffset=False)
        ax.spines[["top", "right"]].set_visible(False)
    fig.suptitle(f"Gene / transcript structure · {detail['protein_accession']}", fontsize=13, weight="bold")
    fig.text(0.5, 0.015, "Teal: CDS · green: annotated exon · yellow: UTR · grey: transcript/intron span. Missing features are not invented.", ha="center", fontsize=8)
    fig.tight_layout(rect=(0, 0.06, 1, 0.94))
    output = io.BytesIO()
    fig.savefig(output, format="png", dpi=160)
    plt.close(fig)
    return output.getvalue()


def bundle(queries, jobs, details):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("report.xlsx", workbook(queries, jobs, details))
        for name, body in fasta_files(details).items():
            if body:
                archive.writestr(name, body)
        archive.writestr("details.json", json.dumps(details, indent=2))
        for index, (_, detail) in enumerate(details.items(), 1):
            archive.writestr(f"gene_structure_{index}.png", structure_figure(detail))
        for i, job in enumerate(jobs, 1):
            if "raw_blast" in job:
                archive.writestr(f"blast_query_{i}.json", json.dumps(job["raw_blast"], indent=2))
    return output.getvalue()
