import io
import zipfile
import json
from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from explorer.demo import demo_data
from explorer.reports import bundle, workbook, structure_figure
from explorer.ncbi import NCBI

APP = Path(__file__).resolve().parents[1] / "app.py"


def test_excel_and_full_download_bundle():
    queries, jobs, details = demo_data()
    wb = load_workbook(io.BytesIO(workbook(queries, jobs, details)))
    assert "Translation_Checks" in wb.sheetnames
    assert wb["Candidates"].max_row == 2
    assert wb["Queries"]["A2"].value == 1
    assert wb["Gene_Loci"].freeze_panes == "A2"
    with zipfile.ZipFile(io.BytesIO(bundle(queries, jobs, details))) as archive:
        assert {"report.xlsx", "cds.fasta", "transcripts.fasta", "genomic_dna.fasta", "details.json", "gene_structure_1.png"} <= set(archive.namelist())
        assert "ATGGCTGAATTTGGTTAA" in archive.read("cds.fasta").decode()
    assert structure_figure(next(iter(details.values()))).startswith(b"\x89PNG")


def test_app_starts_and_offline_demo_shows_details():
    app = AppTest.from_file(APP, default_timeout=30).run()
    assert not app.exception
    next(b for b in app.button if b.label == "Open offline demonstration").click().run()
    assert not app.exception
    assert len(app.tabs) == 5
    assert any("SYNTHETIC" in w.value for w in app.warning)
    assert any(m.value == "DEMO_PROTEIN.1" for m in app.metric)


def test_app_rejects_bad_input_before_submission():
    app = AppTest.from_file(APP, default_timeout=30).run()
    app.text_area(key="protein_text").set_value("MA123EF").run()
    assert any("invalid protein" in e.value for e in app.error)
    assert not app.exception


def test_app_keeps_multiple_query_ids():
    app = AppTest.from_file(APP, default_timeout=30).run()
    app.text_area(key="protein_text").set_value(">same\nMAEFG\n>same\nMLP").run()
    assert not app.exception
    assert app.multiselect[0].value == [0, 1]
    assert any("Duplicate IDs" in info.value for info in app.info)


def test_app_submission_poll_and_candidates(tmp_path, monkeypatch):
    class FixtureNCBI(NCBI):
        def __init__(self, cache_dir, email=""):
            super().__init__(tmp_path, email)

        def request(self, endpoint, params, **kwargs):
            if params.get("CMD") == "Put":
                return "RID = UI123"
            if params.get("CMD") == "Get":
                return json.dumps({"BlastOutput2": [{"report": {"results": {"search": {"query_len": 5, "hits": [{"len": 5, "description": [{"accession": "NP_FIXTURE.1", "sciname": "Synthetic organism", "taxid": 1}], "hsps": [{"query_from": 1, "query_to": 5, "hit_from": 1, "hit_to": 5, "identity": 5, "align_len": 5, "bit_score": 10, "evalue": 1e-9, "qseq": "MAEFG", "hseq": "MAEFG", "midline": "MAEFG"}]}]}}}}]})
            raise AssertionError("Unexpected network call")

    monkeypatch.setattr("explorer.ncbi.NCBI", FixtureNCBI)
    app = AppTest.from_file(APP, default_timeout=30).run()
    app.text_area(key="protein_text").set_value("MAEFG").run()
    next(b for b in app.button if b.label.startswith("Analyse")).click().run()
    assert not app.exception
    assert app.session_state["jobs"][0]["status"] == "WAITING"
    app.session_state["jobs"][0]["last_poll"] = 0
    next(b for b in app.button if b.label == "Check search results").click().run()
    assert not app.exception
    assert app.session_state["jobs"][0]["status"] == "READY"
    candidate = app.session_state["jobs"][0]["candidates"][0]
    assert candidate["accession"] == "NP_FIXTURE.1"
    assert candidate["verification"].startswith("Not verified")
